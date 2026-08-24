import hashlib
import io
import logging
import re
import uuid
from datetime import datetime, timezone

from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.llm.client import LLMClient
from app.models import Chunk, Document
from app.services.usage import record_usage

logger = logging.getLogger(__name__)

# ~3 ký tự/token với tiếng Việt -> khoảng 800 token mỗi đoạn, chồng lấn ~100 token.
TARGET_CHARS = 2400
OVERLAP_CHARS = 300
MIN_CHARS = 200
EMBED_BATCH = 64

SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text",
    "text/markdown": "text",
    "text/x-markdown": "text",
}
EXT_FALLBACK = {".pdf": "pdf", ".docx": "docx", ".md": "text", ".txt": "text"}


class UnsupportedDocument(ValueError):
    pass


def checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def detect_kind(filename: str, content_type: str | None) -> str:
    if content_type and content_type.split(";")[0].strip() in SUPPORTED:
        return SUPPORTED[content_type.split(";")[0].strip()]
    for ext, kind in EXT_FALLBACK.items():
        if filename.lower().endswith(ext):
            return kind
    raise UnsupportedDocument(f"Chỉ hỗ trợ PDF, DOCX, MD, TXT — nhận được '{filename}'")


def extract_text(data: bytes, kind: str) -> str:
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages)
    if kind == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return "\n\n".join(b for b in blocks if b.strip())
    return data.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def split_chunks(text: str) -> list[tuple[str | None, str]]:
    """
    Cắt theo đoạn văn, gom tới ~TARGET_CHARS rồi chồng lấn một phần sang đoạn sau.
    Giữ tiêu đề Markdown gần nhất làm ngữ cảnh để câu trả lời trích dẫn được rõ hơn.
    """
    text = _normalize(text)
    if not text:
        return []

    chunks: list[tuple[str | None, str]] = []
    heading: str | None = None
    buffer = ""

    def flush() -> None:
        nonlocal buffer
        body = buffer.strip()
        if len(body) >= MIN_CHARS or (body and not chunks):
            chunks.append((heading, body))
        buffer = body[-OVERLAP_CHARS:] + "\n\n" if len(body) > OVERLAP_CHARS else ""

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if match := re.match(r"^#{1,6}\s+(.{1,200})$", block):
            if len(buffer.strip()) >= MIN_CHARS:
                flush()
            heading = match.group(1).strip()
            continue
        if len(buffer) + len(block) > TARGET_CHARS and buffer.strip():
            flush()
        buffer += block + "\n\n"

    if buffer.strip():
        chunks.append((heading, buffer.strip()))
    return chunks


async def index_document(
    session: AsyncSession,
    llm: LLMClient,
    settings: Settings,
    document: Document,
    raw_text: str,
) -> int:
    """
    Cắt đoạn, nhúng vector theo lô rồi lưu. Nhúng theo lô là bắt buộc: gọi từng
    đoạn một sẽ chạm trần 50 request/phút chỉ sau vài chục đoạn.
    """
    document.status = "processing"
    document.error = None
    await session.commit()

    try:
        pieces = split_chunks(raw_text)
        if not pieces:
            raise ValueError("Không trích được nội dung văn bản nào từ tài liệu")

        # Index lại từ đầu -> xóa đoạn cũ.
        await session.execute(delete(Chunk).where(Chunk.document_id == document.id))

        total_tokens = 0
        for start in range(0, len(pieces), EMBED_BATCH):
            batch = pieces[start : start + EMBED_BATCH]
            payloads = [
                f"{heading}\n\n{body}" if heading else body for heading, body in batch
            ]
            vectors, tokens = await llm.embed(model=settings.model_embedding, inputs=payloads)
            total_tokens += tokens
            if len(vectors) != len(batch):
                raise ValueError(
                    f"Số vector trả về ({len(vectors)}) khác số đoạn gửi đi ({len(batch)})"
                )
            for offset, ((heading, body), vector) in enumerate(zip(batch, vectors, strict=True)):
                if len(vector) != settings.embedding_dim:
                    raise ValueError(
                        f"Model trả vector {len(vector)} chiều nhưng cột cấu hình "
                        f"{settings.embedding_dim} — sửa EMBEDDING_DIM rồi migrate lại"
                    )
                session.add(
                    Chunk(
                        id=uuid.uuid4(),
                        document_id=document.id,
                        ordinal=start + offset,
                        heading=heading,
                        content=body,
                        token_estimate=len(body) // 3,
                        embedding=vector,
                    )
                )
            await session.commit()

        document.chunk_count = len(pieces)
        document.status = "indexed"
        document.indexed_at = datetime.now(timezone.utc)
        await session.commit()

        await record_usage(
            session,
            settings,
            feature="embedding",
            model=settings.model_embedding,
            input_tokens=total_tokens,
            output_tokens=0,
        )
        logger.info("Đã index tài liệu %s: %d đoạn", document.id, len(pieces))
        return len(pieces)

    except Exception as exc:  # noqa: BLE001 - trạng thái lỗi phải nhìn thấy được qua API
        await session.rollback()
        document.status = "failed"
        document.error = str(exc)[:1000]
        await session.commit()
        logger.exception("Index tài liệu %s thất bại", document.id)
        raise
